"""Excel output for one reconciliation run."""
from __future__ import annotations

import json
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .engine import open_items_as_of, run_items

INR = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
FONT = "Arial"
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=10, bold=True, color="FFFFFF")
FT = Font(name=FONT, size=14, bold=True)
FS = Font(name=FONT, size=9, italic=True, color="555555")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
SUB_FILL = PatternFill("solid", fgColor="E8EEF5")
SEV_FILL = {"critical": "F8D7DA", "high": "FDE8D0", "medium": "FFF4CC", "info": "E8EEF5"}
THIN = Side(style="thin", color="BBBBBB")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
TOP2 = Border(top=Side(style="thin", color="000000"), bottom=Side(style="double", color="000000"))

CAT = {
    ("book", "out"): "A · Payments in book, not yet debited by bank",
    ("book", "in"): "B · Receipts in book, not yet credited by bank",
    ("bank", "in"): "C · Credits in bank, not entered in book",
    ("bank", "out"): "D · Debits in bank, not entered in book",
}


def _table(ws, row, headers, widths=None):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.border = FH, HEAD_FILL, BOX
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    return row + 1


def _put(ws, row, values, amount_cols=(), date_cols=(), fill=None):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.font, cell.border = F, BOX
        if c in amount_cols:
            cell.number_format = INR
        if c in date_cols:
            cell.number_format = "DD-MM-YYYY"
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="top", wrap_text=c not in amount_cols and c not in date_cols)


def _review_text(flag) -> str:
    rv = (flag or {}).get("review")
    if not rv:
        return ""
    return f"{rv['label']} — {rv['username']} {rv['at'][:16]}" + (f": {rv['remark']}" if rv.get("remark") else "")


def build_report(conn, run_id: int) -> tuple[bytes, str]:
    run = conn.execute("SELECT * FROM runs WHERE id=?", (run_id,)).fetchone()
    acc = conn.execute("SELECT * FROM accounts WHERE id=?", (run["account_id"],)).fetchone()
    checks = json.loads(run["checks_json"] or "[]")
    flags = json.loads(run["flags_json"] or "[]")
    brs = json.loads(run["brs_json"] or "{}")
    summ = json.loads(run["summary_json"] or "{}")
    items = run_items(conn, run_id)
    open_ = open_items_as_of(conn, run["account_id"], run_id)
    matches = {m["id"]: dict(m) for m in conn.execute("SELECT * FROM matches WHERE run_id=?", (run_id,))}
    match_run = {m["id"]: m["run_id"] for m in conn.execute(
        "SELECT id, run_id FROM matches WHERE account_id=?", (run["account_id"],))}

    def status(it):
        mid = it.get("match_id")
        if not mid:
            return "OPEN", True
        if match_run.get(mid, 0) > run_id:
            return f"OPEN (cleared later, run #{match_run[mid]})", True
        return f"Matched #{mid}", False
    as_of = date.fromisoformat(run["period_to"])
    p_from = date.fromisoformat(run["period_from"])
    period = f"{p_from:%d-%m-%Y}" + (f" to {as_of:%d-%m-%Y}" if as_of != p_from else "")
    title_acc = f"{acc['bank']} · A/c …{acc['acc_no'][-4:]} · {acc['nickname']}"

    wb = Workbook()

    # ------------------------------------------------------------------ Open items (built first; BRS refers to it)
    wo = wb.active
    wo.title = "Open Items"
    wo["A1"], wo["A1"].font = f"Open (unreconciled) items as at {as_of:%d-%m-%Y} — {title_acc}", FB
    from .engine import holiday_dates, attach_reviews
    hols = holiday_dates(conn, run["account_id"], run_id)
    attach_reviews(conn, run, flags)
    sev_rank = {"critical": 0, "high": 1, "medium": 2, "info": 3}
    worst: dict[str, dict] = {}
    for fl in flags:
        for k in fl.get("keys") or []:
            if k not in worst or sev_rank[fl["severity"]] < sev_rank[worst[k]["severity"]]:
                worst[k] = fl
    r = _table(wo, 3, ["Category", "Side", "Date", "Days open", "Vchr / Sl", "Name", "Details", "Amount",
                       "Carried from run", "Note", "Holiday", "Exception level", "Exception", "Review"],
               [46, 7, 12, 10, 12, 28, 60, 16, 12, 30, 34, 14, 46, 40])
    first = r
    for it in sorted(open_, key=lambda i: (CAT[(i["side"], i["dir"])], i["date"])):
        _put(wo, r, [CAT[(it["side"], it["dir"])], it["side"].title(), it["date"], (as_of - it["date"]).days,
                     it.get("vno") or it.get("sl", ""), it["name"],
                     it.get("narration") or it.get("text", ""), it["amount"],
                     f"#{it['run_id']}" if it["run_id"] != run_id else "", it.get("note", ""),
                     (f"Holiday: received {it['date']:%a %d-%m-%Y}" if it["side"] == "bank" and it["date"] in hols else ""),
                     (worst[it["key"]]["severity"].upper() if it["key"] in worst else ""),
                     (worst[it["key"]]["title"] if it["key"] in worst else ""),
                     _review_text(worst.get(it["key"]))],
             amount_cols=(8,), date_cols=(3,))
        if it["key"] in worst:
            wo.cell(row=r, column=12).fill = PatternFill("solid", fgColor=SEV_FILL[worst[it["key"]]["severity"]])
        r += 1
    last = max(r - 1, first)
    if not open_:
        _put(wo, r, ["No open items — every line is matched."] + [""] * 13)
        last = r
    cat_rng = f"'Open Items'!$A${first}:$A${last}"
    amt_rng = f"'Open Items'!$H${first}:$H${last}"
    hol_rng = f"'Open Items'!$K${first}:$K${last}"

    # ------------------------------------------------------------------ BRS (first tab)
    ws = wb.create_sheet("BRS", 0)
    ws.column_dimensions["A"].width = 62
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 20
    ws["A1"], ws["A1"].font = "Bank Reconciliation Statement", FT
    ws["A2"], ws["A2"].font = title_acc, FB
    ws["A3"], ws["A3"].font = f"As at {as_of:%d-%m-%Y} (period {period}) · Run #{run_id}", F
    ws["C5"], ws["C5"].font = "Amount (₹)", FB
    ws["B5"], ws["B5"].font = "Items", FB
    lines = [
        (7, "Balance as per bank book", None, run["book_close"], None),
        (8, "Add: Payments entered in book, not yet debited by bank", "A", None, "a"),
        (9, "Less: Receipts entered in book, not yet credited by bank", "B", None, "b"),
        (10, "Add: Credits in bank statement, not entered in book", "C", None, "c"),
        (11, "Less: Debits in bank statement, not entered in book", "D", None, "d"),
        (12, "Add/(Less): Opening difference carried, not itemised", None, brs.get("opening_diff", 0), None),
    ]
    for row, label, cat, val, cnt in lines:
        ws.cell(row=row, column=1, value=label).font = FB if row == 7 else F
        c = ws.cell(row=row, column=3)
        if cat:
            c.value = f'=SUMIFS({amt_rng},{cat_rng},"{cat} ·*")'
            ws.cell(row=row, column=2, value=f'=COUNTIFS({cat_rng},"{cat} ·*")').font = F
        else:
            c.value = val
        c.number_format, c.font = INR, F
    ws["A13"], ws["A13"].font = "      of which received on a holiday (due in next working day's book) — memo", FS
    ws["B13"] = f'=COUNTIFS({cat_rng},"C ·*",{hol_rng},"Holiday*")'
    ws["C13"] = f'=SUMIFS({amt_rng},{cat_rng},"C ·*",{hol_rng},"Holiday*")'
    ws["B13"].font = ws["C13"].font = FS
    ws["C13"].number_format = INR
    ws["A14"], ws["A14"].font = "Balance as per bank statement (computed)", FB
    ws["C14"] = "=C7+C8-C9+C10-C11+C12"
    ws["A15"], ws["A15"].font = "Balance as per bank statement (actual)", FB
    ws["C15"] = run["bank_close"]
    ws["A16"], ws["A16"].font = "Difference (must be zero)", FB
    ws["C16"] = "=ROUND(C15-C14,2)"
    for a in ("C14", "C15", "C16"):
        ws[a].number_format, ws[a].font = INR, FB
    ws["C14"].border = TOP2
    ws["A17"] = ('=IF(ABS(C16)>=0.005,"✖ NOT RECONCILED — investigate the difference",'
                 'IF(ABS(C12)>=0.005,"⚠ Balances only with an UNEXPLAINED difference (row 12) — not a clean reconciliation",'
                 '"✔ Reconciled — difference is nil"))')
    ws["A17"].font = Font(name=FONT, size=11, bold=True, color="1B7F3B" if brs.get("clean") else "B00020")
    ws["A19"], ws["A19"].font = "Source figures", FB
    src = [("Bank statement opening", run["bank_open"]), ("Bank statement closing", run["bank_close"]),
           ("Bank book opening", run["book_open"]), ("Bank book closing", run["book_close"])]
    for i, (k, v) in enumerate(src, 20):
        ws.cell(row=i, column=1, value=k).font = F
        c = ws.cell(row=i, column=3, value=v)
        c.number_format, c.font = INR, F
    ws["A25"], ws["A25"].font = "Files", FB
    ws["A26"], ws["A26"].font = f"Bank statement: {run['bank_file']}", F
    ws["A27"], ws["A27"].font = f"Bank book: {run['book_file']}", F
    ws["A28"], ws["A28"].font = f"Prepared by system for {run['created_by']} on {run['created_at']}", FS
    if run["ack_reason"]:
        ws["A29"], ws["A29"].font = f"Warnings acknowledged: {run['ack_reason']}", FS
    ws["A31"], ws["A31"].font = f"Prepared by: {run['created_by']}  ({run['created_at']})", F
    chk = run["checked_by"] if "checked_by" in run.keys() else ""
    ws["A33"], ws["A33"].font = (f"Checked by: {chk}  ({run['checked_at']})" + (f" — {run['checked_note']}" if run["checked_note"] else "")
                                 if chk else "Checked by: ____________________"), F
    ws["A35"], ws["A35"].font = "Approved by: ____________________", F

    # ------------------------------------------------------------------ Summary & checks
    wsum = wb.create_sheet("Summary & Checks", 1)
    wsum["A1"], wsum["A1"].font = "Summary", FT
    r = _table(wsum, 3, ["", "Bank lines", "Bank amount", "Matched", "Book entries", "Book amount", "Matched"],
               [34, 12, 18, 12, 14, 18, 12])
    for label, bk, kk in (("Payments / debits", "bank_out", "book_out"), ("Collections / credits", "bank_in", "book_in")):
        b, k = summ.get(bk, {}), summ.get(kk, {})
        _put(wsum, r, [label, b.get("count"), b.get("amount"), b.get("matched"),
                       k.get("count"), k.get("amount"), k.get("matched")], amount_cols=(3, 6))
        r += 1
    r += 1
    for label, v in (("Matches made", summ.get("matches")), ("…of which need review", summ.get("review_matches")),
                     ("…of which manual", summ.get("manual_matches")),
                     ("Earlier open items cleared in this run", summ.get("carried_cleared")),
                     ("Open items after this run", summ.get("open_items"))):
        _put(wsum, r, [label, v])
        r += 1
    r += 1
    wsum.cell(row=r, column=1, value="Validation checks").font = FT
    r = _table(wsum, r + 1, ["Result", "Source", "Check", "Detail"])
    wsum.column_dimensions["D"].width = 100
    for c in checks:
        fill = {"ok": "E3F4E8", "warning": "FFF4CC", "error": "F8D7DA"}[c["level"]]
        _put(wsum, r, [{"ok": "PASS", "warning": "WARN", "error": "FAIL"}[c["level"]], c.get("source", ""),
                       c["code"], c["message"]], fill=fill)
        r += 1
    wsum.freeze_panes = None

    # ------------------------------------------------------------------ Exceptions
    wx = wb.create_sheet("Exceptions", 2)
    r = _table(wx, 1, ["Severity", "Type", "Title", "Detail", "Review status", "Review remark", "Reviewed by", "Reviewed at"],
               [11, 26, 55, 110, 18, 50, 16, 20])
    for f in flags:
        rv = f.get("review") or {}
        _put(wx, r, [f["severity"].upper(), f["code"], f["title"], f["detail"], rv.get("label", "Not reviewed"),
                     rv.get("remark", ""), rv.get("username", ""), rv.get("at", "")], fill=SEV_FILL[f["severity"]])
        r += 1
    if not flags:
        _put(wx, r, ["", "", "No exceptions.", ""])

    # ------------------------------------------------------------------ matched detail
    byid = {}
    for it in items:
        if it.get("match_id") in matches:
            byid.setdefault(it["match_id"], []).append(it)

    def matched_sheet(title, direction):
        wm = wb.create_sheet(title)
        r = _table(wm, 1, ["Match #", "Method", "Confidence", "Side", "Date", "Sl / Vchr", "Name",
                           "Reference / Loan A/c", "Amount", "Note"], [9, 34, 11, 7, 12, 10, 28, 26, 16, 60])
        for mid in sorted(byid):
            its = [i for i in byid[mid] if i["dir"] == direction]
            if not its:
                continue
            m = matches[mid]
            fill = None if m["confidence"] == "strong" else ("FFF4CC" if m["confidence"] == "review" else "E8EEF5")
            for it in sorted(its, key=lambda i: i["side"]):
                ref = it.get("ref") if it["side"] == "bank" else (it.get("loan_ac") or it.get("ref", ""))
                _put(wm, r, [mid, m["method"], m["confidence"], it["side"].title(), it["date"],
                             it.get("sl") if it["side"] == "bank" else it.get("vno"), it["name"], ref,
                             it["amount"], m["note"] + (" (manual, by " + m["created_by"] + ")" if m["manual"] else "")],
                     amount_cols=(9,), date_cols=(5,), fill=fill)
                r += 1

    matched_sheet("Payments Matched", "out")
    matched_sheet("Collections Matched", "in")

    # ------------------------------------------------------------------ source data with status
    own = [i for i in items if i["run_id"] == run_id]
    wbk = wb.create_sheet("Bank Statement")
    r = _table(wbk, 1, ["Sl", "Date", "Particulars", "Customer name", "Withdrawal", "Deposit", "Balance", "Channel",
                        "UTR / Ref", "Status"], [6, 12, 70, 28, 15, 15, 18, 8, 16, 18])
    for it in sorted((i for i in own if i["side"] == "bank"), key=lambda i: i["row_no"]):
        st, is_open = status(it)
        _put(wbk, r, [it.get("sl"), it["date"], it["text"], it.get("name") or "",
                      it["amount"] if it["dir"] == "out" else None,
                      it["amount"] if it["dir"] == "in" else None, it.get("balance"), it.get("channel"),
                      it.get("ref"), st], amount_cols=(5, 6, 7), date_cols=(2,),
             fill="F8D7DA" if is_open else None)
        r += 1
    wbb = wb.create_sheet("Bank Book")
    r = _table(wbb, 1, ["Vchr", "Type", "Branch", "Date", "Name", "Loan A/c", "Instrument", "Narration",
                        "Receipt (Dr)", "Payment (Cr)", "Status"], [9, 6, 16, 12, 26, 22, 22, 60, 15, 15, 16])
    for it in sorted((i for i in own if i["side"] == "book"), key=lambda i: i["row_no"]):
        st, is_open = status(it)
        _put(wbb, r, [it.get("vno"), it.get("vtype"), it.get("branch"), it["date"], it["name"], it.get("loan_ac"),
                      it.get("ref"), it.get("narration"), it["amount"] if it["dir"] == "in" else None,
                      it["amount"] if it["dir"] == "out" else None, st], amount_cols=(9, 10), date_cols=(4,),
             fill="F8D7DA" if is_open else None)
        r += 1

    for sh in wb.worksheets:
        sh.sheet_view.showGridLines = sh.title != "BRS"
        sh.page_setup.orientation = "landscape"
        sh.page_setup.fitToWidth = 1
    ws.sheet_view.showGridLines = False

    buf = BytesIO()
    wb.save(buf)
    fname = f"BRS_{acc['bank'].replace(' ', '')}_{acc['acc_no'][-4:]}_{as_of:%d-%m-%Y}.xlsx"
    return buf.getvalue(), fname
