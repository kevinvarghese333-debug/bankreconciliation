"""Bulk import of bank accounts from Excel/CSV.

Two layouts are accepted:
1. The template (header row): Bank | Account number | Nickname | Type | Branch | Notes | Active
2. Free-form lists, e.g.  "Aloor | SOUTH INDIAN BANK A/C NO 0058073000001317".
   The longest run of 8+ digits in a row is the account number, a bank name is recognised from the text,
   and the remaining text cell is taken as the branch.
"""
from __future__ import annotations

import re
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import db
from .parsers import norm, read_sheet

BANKS = {
    "southindianbank": "South Indian Bank", "sib": "South Indian Bank",
    "federalbank": "Federal Bank", "hdfc": "HDFC Bank", "icici": "ICICI Bank", "axis": "Axis Bank",
    "statebankofindia": "State Bank of India", "sbi": "State Bank of India", "canarabank": "Canara Bank",
    "unionbank": "Union Bank of India", "bankofbaroda": "Bank of Baroda", "kotak": "Kotak Mahindra Bank",
    "idbi": "IDBI Bank", "csb": "CSB Bank", "dhanlaxmi": "Dhanlaxmi Bank", "esafsmallfinance": "ESAF Small Finance Bank",
}
TEMPLATE_HEADERS = ["Bank", "Account number", "Nickname", "Type", "Branch", "Notes", "Active"]
DIGITS = re.compile(r"(?<!\d)(\d[\d\s-]{6,}\d)(?!\d)")


def _bank_from(text: str) -> str:
    t = norm(text)
    for k, v in sorted(BANKS.items(), key=lambda x: -len(x[0])):
        if k in t:
            return v
    return ""


def _kind(v: str) -> str | None:
    v = norm(v)
    if not v:
        return None
    if v.startswith("main") or v in ("payment", "payments", "disbursement", "disbursements"):
        return "main"
    if v.startswith("coll") or v in ("receipt", "receipts"):
        return "collection"
    return "invalid"


def parse_accounts(path: str, default_kind: str = "collection", default_bank: str = "South Indian Bank",
                   main_last4: set[str] | None = None) -> tuple[list[dict], list[str]]:
    """Returns (rows, problems). Each row: bank, acc_no, nickname, kind, branch, notes, active, source_row."""
    main_last4 = main_last4 or set()
    df = read_sheet(path)
    rows, problems = [], []
    header_idx, cols = None, {}
    for i in range(min(15, len(df))):
        cells = [norm(c) if not pd.isna(c) else "" for c in df.iloc[i].tolist()]
        if any(c in ("accountnumber", "accountno", "acno", "acnumber", "account") for c in cells):
            header_idx = i
            for j, c in enumerate(cells):
                if c in ("bank", "bankname"): cols["bank"] = j
                elif c in ("accountnumber", "accountno", "acno", "acnumber", "account"): cols["acc"] = j
                elif c in ("nickname", "name", "accountname"): cols["nick"] = j
                elif c in ("type", "kind", "accounttype"): cols["kind"] = j
                elif c in ("branch", "branchname"): cols["branch"] = j
                elif c in ("notes", "note", "remarks"): cols["notes"] = j
                elif c in ("active", "status"): cols["active"] = j
            break

    start = header_idx + 1 if header_idx is not None else 0
    for i in range(start, len(df)):
        raw = ["" if pd.isna(c) else str(c).strip() for c in df.iloc[i].tolist()]
        if not any(raw):
            continue
        excel_row = i + 1

        def cell(k):
            j = cols.get(k)
            return raw[j] if j is not None and j < len(raw) else ""

        if header_idx is not None and "acc" in cols:
            acc_text = cell("acc")
            if acc_text.endswith(".0"):
                acc_text = acc_text[:-2]
            acc = re.sub(r"\D", "", acc_text)
            bank = cell("bank") or default_bank
            bank = _bank_from(bank) or bank
            branch, nick, notes = cell("branch"), cell("nick"), cell("notes")
            k = _kind(cell("kind"))
            if k == "invalid":
                problems.append(f"Row {excel_row}: Type '{cell('kind')}' not understood — use Main or Collection.")
                continue
            act = norm(cell("active"))
            active = 0 if act in ("no", "n", "0", "inactive", "false") else 1
        else:  # free-form row
            joined = " ".join(raw)
            m = max(DIGITS.findall(joined), key=lambda s: len(re.sub(r"\D", "", s)), default="")
            acc = re.sub(r"\D", "", m)
            bank = _bank_from(joined) or default_bank
            others = [c for c in raw if c and not DIGITS.search(c)]
            branch = others[0] if others else ""
            nick, notes, k, active = "", "", None, 1
        if not acc:
            problems.append(f"Row {excel_row}: no account number found — skipped ({' | '.join(c for c in raw if c)[:80]}).")
            continue
        if len(acc) < 4:
            problems.append(f"Row {excel_row}: account number '{acc}' too short — skipped.")
            continue
        kind = k or ("main" if acc[-4:] in main_last4 else default_kind)
        if not nick:
            short = "SIB" if bank == "South Indian Bank" else bank
            nick = f"{short} {acc[-4:]} · {branch}" if branch else f"{short} {acc[-4:]}"
        rows.append({"bank": bank, "acc_no": acc, "nickname": nick[:80], "kind": kind, "branch": branch,
                     "notes": notes, "active": active, "source_row": excel_row})

    # duplicates inside the file
    seen = {}
    for r in rows:
        key = (r["bank"], r["acc_no"])
        if key in seen:
            problems.append(f"Row {r['source_row']}: account {r['acc_no']} repeats row {seen[key]} — later row used.")
        seen[key] = r["source_row"]
    # unusual length compared with the rest of the file
    lengths = {}
    for r in rows:
        lengths.setdefault(r["bank"], []).append(len(r["acc_no"]))
    for r in rows:
        ls = lengths[r["bank"]]
        common = max(set(ls), key=ls.count)
        if len(ls) >= 3 and len(r["acc_no"]) != common:
            problems.append(f"Row {r['source_row']}: {r['branch'] or r['nickname']} account {r['acc_no']} has "
                            f"{len(r['acc_no'])} digits; the others have {common}. Imported — please check it.")
    return rows, problems


def apply_import(conn, rows: list[dict], user: str) -> dict:
    added = updated = unchanged = 0
    notes = []
    with db.tx(conn):
        for r in {(x["bank"], x["acc_no"]): x for x in rows}.values():
            old = conn.execute("SELECT * FROM accounts WHERE bank=? AND acc_no=?", (r["bank"], r["acc_no"])).fetchone()
            if not old:
                conn.execute("INSERT INTO accounts(bank, acc_no, nickname, kind, branch, notes, active, created_at) "
                             "VALUES (?,?,?,?,?,?,?,?)",
                             (r["bank"], r["acc_no"], r["nickname"], r["kind"], r["branch"], r["notes"],
                              r["active"], db.now()))
                added += 1
                continue
            used = conn.execute("SELECT COUNT(*) FROM runs WHERE account_id=?", (old["id"],)).fetchone()[0]
            kind = r["kind"]
            if used and kind != old["kind"]:
                notes.append(f"{old['nickname']}: type kept as {old['kind']} (already reconciled).")
                kind = old["kind"]
            new = (r["nickname"], kind, r["branch"], r["notes"] or old["notes"], r["active"])
            cur = (old["nickname"], old["kind"], old["branch"], old["notes"], old["active"])
            if new == cur:
                unchanged += 1
                continue
            conn.execute("UPDATE accounts SET nickname=?, kind=?, branch=?, notes=?, active=? WHERE id=?",
                         (*new, old["id"]))
            updated += 1
        db.audit(conn, user, "accounts_import", f"{added} added, {updated} updated, {unchanged} unchanged")
    return {"added": added, "updated": updated, "unchanged": unchanged, "notes": notes}


def template_bytes(example_rows: list[list] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    ws.append(TEMPLATE_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    for r in example_rows or [["South Indian Bank", "0058073000001113", "SIB 1113 · North Paravur (main)", "Main",
                               "North Paravur", "Disbursement account", "Yes"]]:
        ws.append(r)
    for col in ws.iter_cols(min_col=2, max_col=2, min_row=2, max_row=500):
        for c in col:
            c.number_format = "@"  # keep leading zeros
    dv = DataValidation(type="list", formula1='"Main,Collection"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("D2:D500")
    dv2 = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv2)
    dv2.add("G2:G500")
    for col, w in zip("ABCDEFG", (20, 22, 36, 12, 22, 30, 8)):
        ws.column_dimensions[col].width = w
    help_ = wb.create_sheet("How to fill")
    for line in ["One row per bank account. Only 'Account number' is required.",
                 "Type Account number as text so leading zeros stay (this sheet is already formatted that way).",
                 "Type: Main = payments/disbursements account; Collection = receipts + day-end sweep. Blank = Collection.",
                 "Nickname: leave blank to auto-name as 'SIB 1317 · Branch'.",
                 "Uploading again updates existing accounts (matched by bank + account number); nothing is deleted.",
                 "Active: No hides the account from new reconciliations."]:
        help_.append([line])
    help_.column_dimensions["A"].width = 110
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
