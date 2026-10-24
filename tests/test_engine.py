"""End-to-end tests. Run:  python -m pytest -q   (from the project folder)"""
import json
import os
import shutil
import sys
import tempfile
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
TMP = tempfile.mkdtemp(prefix="bankreco-test-")
os.environ["DATA_DIR"] = TMP

from app import db, engine  # noqa: E402
from tests.synth import write_book, write_statement  # noqa: E402

SAMPLES = ROOT / "samples"


@pytest.fixture()
def conn():
    shutil.rmtree(TMP, ignore_errors=True)
    db.init()
    c = db.connect()
    c.execute("INSERT INTO accounts(bank,acc_no,nickname,kind,created_at) VALUES('SIB','1113','Main','main',?)", (db.now(),))
    c.execute("INSERT INTO accounts(bank,acc_no,nickname,kind,created_at) VALUES('SIB','2201','Coll','collection',?)", (db.now(),))
    yield c
    c.close()


def run(conn, acc, bank, book, **kw):
    """Reconcile; a first run's mandatory account confirmation is acknowledged automatically."""
    try:
        return engine.reconcile(conn, acc, str(bank), Path(bank).name, str(book), Path(book).name, "test", **kw)
    except engine.Blocked as e:
        if e.ackable and {c["code"] for c in e.ackable} == {"FIRST_RUN"} and not kw.get("acknowledged"):
            return engine.reconcile(conn, acc, str(bank), Path(bank).name, str(book), Path(book).name, "test",
                                    acknowledged=True, ack_reason="Confirmed account for first run")
        raise


def brs(conn, rid):
    return json.loads(conn.execute("SELECT brs_json FROM runs WHERE id=?", (rid,)).fetchone()[0])


def flags(conn, rid):
    return json.loads(conn.execute("SELECT flags_json FROM runs WHERE id=?", (rid,)).fetchone()[0])


def codes(conn, rid):
    return [f["code"] for f in flags(conn, rid)]


def test_real_sample_tallies(conn):
    if not (SAMPLES / "SIB_1113_26-09-2026.xls").exists():
        pytest.skip("sample files not present")
    rid = run(conn, 1, SAMPLES / "SIB_1113_26-09-2026.xls", SAMPLES / "BankBook_1113_26-09-2026.xlsx")
    b = brs(conn, rid)
    assert b["tallies"] and b["difference"] == 0
    s = json.loads(conn.execute("SELECT summary_json FROM runs WHERE id=?", (rid,)).fetchone()[0])
    assert s["bank_out"]["matched"] == s["bank_out"]["count"] == 109
    assert "BANK_DEBIT_NOT_IN_BOOK" not in codes(conn, rid)


def test_carry_forward_and_exceptions(conn, tmp_path):
    d1, d2 = date(2026, 9, 28), date(2026, 9, 29)
    op = 1_000_000.00
    # Day 1: A paid & debited; B booked but bank debits tomorrow; bank charge not booked; UPI collection
    bank1 = tmp_path / "b1.xlsx"
    bclose1 = write_statement(bank1, d1, op, [
        ("RAP/626900000001/ 0000095001 /ANITHA KUMARI/", 50000, 0),
        ("UPI/SBIN/626911112222/RAJI BIJEESH/UPI/x/raji@ok", 0, 800),
        ("CHRG/SMS CHARGES Q3", 59, 0),
    ])
    book1 = tmp_path / "k1.xlsx"
    kclose1 = write_book(book1, d1, op, [
        dict(vno=1, name="ANITHA KUMARI", loan_ac="0001", payment=50000),
        dict(vno=2, name="BINDU RAJAN", loan_ac="0002", payment=30000),
        dict(vno=3, name="RAJI BIJEESH", loan_ac="0003", instr="2222", receipt=800),
    ])
    r1 = run(conn, 1, bank1, book1)
    b = brs(conn, r1)
    assert b["tallies"], b
    assert b["add_book_payments_not_debited"] == 30000
    assert b["less_bank_debits_not_in_book"] == 59
    c = codes(conn, r1)
    assert "TIMING_BANK_DEBIT" in c and "TIMING_BOOK_PAYMENT" in c       # first day: timing difference

    # Day 2: B debited by bank (clears carry); charge booked (clears carry); fake receipt in book only
    bank2 = tmp_path / "b2.xlsx"
    write_statement(bank2, d2, bclose1, [("RAP/626900000009/ 0000095002 /BINDU RAJAN/", 30000, 0)])
    book2 = tmp_path / "k2.xlsx"
    write_book(book2, d2, kclose1, [
        dict(vno=4, name="", payment=59, vtype="JV", narration="Bank charges SMS Q3"),
        dict(vno=5, name="SUMA K", loan_ac="0005", instr="9999", receipt=1000),
    ])
    r2 = run(conn, 1, bank2, book2)
    b = brs(conn, r2)
    assert b["tallies"], b
    # BINDU's ₹30,000 cleared. The ₹59 charge booked a day late has no name/ref, so it is NOT
    # auto-cleared against the carried bank line — both sides stay open until matched by hand.
    assert b["add_book_payments_not_debited"] == 59
    assert b["less_bank_debits_not_in_book"] == 59
    open_ = engine.open_items_as_of(conn, 1, r2)
    ids = [i["id"] for i in open_ if i["amount"] == 59]
    engine.manual_match(conn, 1, ids, "Bank SMS charges booked next day", "test")
    b = brs(conn, r2)
    assert b["tallies"] and b["less_bank_debits_not_in_book"] == 0 and b["add_book_payments_not_debited"] == 0
    assert b["less_book_receipts_not_credited"] == 1000
    c = codes(conn, r2)
    assert "TIMING_BOOK_RECEIPT" in c
    assert "BANK_DEBIT_NOT_IN_BOOK" not in c and "TIMING_BANK_DEBIT" not in c
    s = json.loads(conn.execute("SELECT summary_json FROM runs WHERE id=?", (r2,)).fetchone()[0])
    assert s["carried_cleared"] >= 1

    # Undo day 2 -> day-1 items reopen
    engine.undo_run(conn, r2, "test")
    assert brs(conn, r1)["add_book_payments_not_debited"] == 30000


def test_tampered_statement_blocked(conn, tmp_path):
    import openpyxl
    bank = tmp_path / "b.xlsx"
    write_statement(bank, date(2026, 9, 28), 1000.0,
                    [("RAP/1/ 1 /A/", 100, 0), ("RAP/2/ 2 /X/", 50, 0), ("RAP/3/ 3 /B/", 200, 0)])
    wb = openpyxl.load_workbook(bank)
    wb.active.delete_rows(3)  # quietly remove the middle transaction
    wb.save(bank)
    book = tmp_path / "k.xlsx"
    write_book(book, date(2026, 9, 28), 1000.0, [dict(vno=1, name="B", payment=200)])
    with pytest.raises(engine.Blocked) as e:
        run(conn, 1, bank, book)
    assert any(c["code"] == "BANK_BALANCE_BREAK" and c["level"] == "error" for c in e.value.checks)
    assert e.value.ackable == []  # cannot be overridden


def test_continuity_break_needs_ack(conn, tmp_path):
    d1, d2 = date(2026, 9, 28), date(2026, 9, 30)
    b1, k1 = tmp_path / "b1.xlsx", tmp_path / "k1.xlsx"
    write_statement(b1, d1, 5000, [("RAP/1/ 1 /A/", 100, 0)])
    write_book(k1, d1, 5000, [dict(vno=1, name="A", payment=100)])
    run(conn, 1, b1, k1)
    b2, k2 = tmp_path / "b2.xlsx", tmp_path / "k2.xlsx"
    write_statement(b2, d2, 4800, [("RAP/2/ 2 /C/", 50, 0)])   # 100 missing vs previous closing 4900
    write_book(k2, d2, 4900, [dict(vno=2, name="C", payment=50)])
    with pytest.raises(engine.Blocked) as e:
        run(conn, 1, b2, k2, acknowledged=True, ack_reason="Missing statement for 28th, checked with bank")
    assert any(c["code"] == "BANK_CONTINUITY" and c["level"] == "error" for c in e.value.checks)
    assert e.value.ackable == []   # a broken bank chain can never be waved through


def test_book_backdated_entry_needs_reason_and_stays_critical(conn, tmp_path):
    d1, d2 = date(2026, 9, 28), date(2026, 9, 29)
    b1, k1 = tmp_path / "b1.xlsx", tmp_path / "k1.xlsx"
    bc = write_statement(b1, d1, 5000, [("RAP/1/ 1 /ANU/", 100, 0)])
    kc = write_book(k1, d1, 5000, [dict(vno=1, name="ANU", payment=100)])
    run(conn, 1, b1, k1)
    b2, k2 = tmp_path / "b2.xlsx", tmp_path / "k2.xlsx"
    write_statement(b2, d2, bc, [("RAP/2/ 2 /CHITRA/", 50, 0)])
    write_book(k2, d2, kc - 700, [dict(vno=2, name="CHITRA", payment=50)])   # back-dated ₹700 entry in ERP
    with pytest.raises(engine.Blocked) as e:
        run(conn, 1, b2, k2)
    assert "BOOK_CONTINUITY" in {c["code"] for c in e.value.ackable}
    with pytest.raises(engine.Blocked):
        run(conn, 1, b2, k2, acknowledged=True, ack_reason=".")          # too short
    rid = run(conn, 1, b2, k2, acknowledged=True, ack_reason="Back-dated JV 28-09 for ₹700, to be traced")
    b = brs(conn, rid)
    assert b["tallies"] and not b["clean"] and b["opening_diff"] == 700
    assert codes(conn, rid)[0] == "UNEXPLAINED_DIFFERENCE"
    assert flags(conn, rid)[0]["severity"] == "critical"


def test_collection_account_outflows(conn, tmp_path):
    d = date(2026, 9, 28)
    b, k = tmp_path / "b.xlsx", tmp_path / "k.xlsx"
    write_statement(b, d, 0, [
        ("UPI/SBIN/626900001111/SHEEBA JOY/UPI/x/s@ok", 0, 800),
        ("UPI/SBIN/626900002222/NIMISHA/UPI/x/n@ok", 0, 700),
        ("NEFT/SIBLN00001/UNKNOWN PARTY/", 1000, 0),
        ("TRF/SWEEP TO 1113", 500, 0),
    ])
    write_book(k, d, 0, [
        dict(vno=1, name="SHEEBA JOY", instr="1111", receipt=800),
        dict(vno=2, name="NIMISHA BAIJU", instr="2222", receipt=700),
        dict(vno=3, name="", payment=500, vtype="CT", narration="Sweep transfer to SIB 1113"),
    ])
    rid = run(conn, 2, b, k, acknowledged=True, ack_reason="First run for collection account")
    f = flags(conn, rid)
    crit = [x for x in f if x["severity"] == "critical"]
    assert len(crit) == 1 and crit[0]["code"] == "BANK_DEBIT_NOT_IN_BOOK"
    assert "COLLECTION_MULTI_OUT" in [x["code"] for x in f]
    assert brs(conn, rid)["tallies"]


def test_manual_match_and_unmatch(conn, tmp_path):
    d = date(2026, 9, 28)
    b, k = tmp_path / "b.xlsx", tmp_path / "k.xlsx"
    write_statement(b, d, 1000, [("RAP/1/ 1 /TOTALLY DIFFERENT/", 100, 0), ("RAP/2/ 2 /OTHER/", 100, 0)])
    write_book(k, d, 1000, [dict(vno=1, name="PERSON A", payment=100), dict(vno=2, name="PERSON B", payment=100)])
    rid = run(conn, 1, b, k)
    open_ = engine.open_items_as_of(conn, 1, rid)
    assert len(open_) == 4
    bank_id = next(i["id"] for i in open_ if i["side"] == "bank")
    book_id = next(i["id"] for i in open_ if i["side"] == "book")
    with pytest.raises(ValueError):
        engine.manual_match(conn, 1, [bank_id, book_id], "", "test")
    mid = engine.manual_match(conn, 1, [bank_id, book_id], "Confirmed with branch", "test")
    assert len(engine.open_items_as_of(conn, 1, rid)) == 2
    engine.unmatch(conn, mid, "test", "wrong pair")
    assert len(engine.open_items_as_of(conn, 1, rid)) == 4
    assert brs(conn, rid)["tallies"]


def test_duplicate_payment_flag(conn, tmp_path):
    d = date(2026, 9, 28)
    b, k = tmp_path / "b.xlsx", tmp_path / "k.xlsx"
    write_statement(b, d, 100000, [("RAP/1/ 1 /LILLY JOHNSON/", 15188, 0), ("RAP/2/ 2 /LILLY JOHNSON/", 15188, 0)])
    write_book(k, d, 100000, [dict(vno=1, name="LILLY JOHNSON", loan_ac="0034", payment=15188),
                              dict(vno=2, name="LILLY JOHNSON", loan_ac="0034", payment=15188)])
    rid = run(conn, 1, b, k)
    c = codes(conn, rid)
    assert "DUP_LOAN_PAYMENT" in c and "DUP_BANK_PAYMENT" in c


def _day(conn, tmp_path, tag, d, op, bank_rows, book_rows, acc=1, book_op=None, **kw):
    b, k = tmp_path / f"b{tag}.xlsx", tmp_path / f"k{tag}.xlsx"
    bc = write_statement(b, d, op, bank_rows)
    kc = write_book(k, d, op if book_op is None else book_op, book_rows)
    return run(conn, acc, b, k, **kw), bc, kc


def _conf(conn, rid, name):
    """confidence of the match that contains the book entry for `name`"""
    for i in engine.run_items(conn, rid):
        if i["side"] == "book" and i["name"] == name:
            if not i["match_id"]:
                return None
            return conn.execute("SELECT confidence FROM matches WHERE id=?", (i["match_id"],)).fetchone()[0]


def test_different_person_same_amount_is_not_strong(conn, tmp_path):
    rid, _, _ = _day(conn, tmp_path, 1, date(2026, 9, 28), 100000,
                     [("RAP/626900000001/ 0000095001 /SUNIL KUMAR/", 50000, 0)],
                     [dict(vno=1, name="ANIL KUMAR", loan_ac="0001", payment=50000)])
    assert _conf(conn, rid, "ANIL KUMAR") in (None, "review")


def test_tied_candidates_are_flagged_ambiguous(conn, tmp_path):
    rid, _, _ = _day(conn, tmp_path, 1, date(2026, 9, 28), 100000,
                     [("RAP/626900000001/ 0000095001 /THOMAS/", 20000, 0)],
                     [dict(vno=1, name="THOMAS VARGHESE", loan_ac="0001", payment=20000),
                      dict(vno=2, name="THOMAS K J", loan_ac="0002", payment=20000)])
    confs = {_conf(conn, rid, "THOMAS VARGHESE"), _conf(conn, rid, "THOMAS K J")}
    assert "strong" not in confs


def test_reference_fragment_does_not_clear(conn, tmp_path):
    d1 = date(2026, 9, 28)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 100000, [("RAP/1/ 1 /X/", 1, 0)],
                      [dict(vno=1, name="X", payment=1), dict(vno=2, name="BINDU", loan_ac="00030400600000000700", payment=25000)])
    b2, k2 = tmp_path / "b2.xlsx", tmp_path / "k2.xlsx"
    write_statement(b2, date(2026, 9, 29), bc, [("CLG/000700/UNKNOWN PERSON/", 25000, 0)])
    write_book(k2, date(2026, 9, 29), kc, [])
    r2 = run(conn, 1, b2, k2)
    assert _conf(conn, r2, "BINDU") in (None, "review")


def test_carried_payment_not_cleared_by_unrelated_debit(conn, tmp_path):
    r1, bc, kc = _day(conn, tmp_path, 1, date(2026, 9, 28), 100000, [("RAP/1/ 1 /X/", 1, 0)],
                      [dict(vno=1, name="X", payment=1), dict(vno=2, name="BINDU RAJAN", payment=30000)])
    b2, k2 = tmp_path / "b2.xlsx", tmp_path / "k2.xlsx"
    write_statement(b2, date(2026, 9, 29), bc, [("RAP/2/ 2 /SURESH MENON/", 30000, 0),
                                                 ("UPI/SBIN/626900007777/TOM/UPI/x/p@ok", 0, 250)])
    write_book(k2, date(2026, 9, 29), kc, [dict(vno=3, name="TOM", loan_ac="07", instr="7777", receipt=250)])
    r2 = run(conn, 1, b2, k2)
    assert _conf(conn, r2, "BINDU RAJAN") is None
    assert "TIMING_BANK_DEBIT" in codes(conn, r2) and "BOOK_PAYMENT_NOT_IN_BANK" in codes(conn, r2)


def test_double_disbursal_across_days(conn, tmp_path):
    r1, bc, kc = _day(conn, tmp_path, 1, date(2026, 9, 28), 100000, [("RAP/1/ 1 /LILLY JOHNSON/", 15188, 0)],
                      [dict(vno=1, name="LILLY JOHNSON", loan_ac="0034", payment=15188)])
    b2, k2 = tmp_path / "b2.xlsx", tmp_path / "k2.xlsx"
    write_statement(b2, date(2026, 9, 29), bc, [("RAP/2/ 2 /LILLY JOHNSON/", 15188, 0)])
    write_book(k2, date(2026, 9, 29), kc, [dict(vno=2, name="LILLY JOHNSON", loan_ac="0034", payment=15188)])
    r2 = run(conn, 1, b2, k2)
    c = codes(conn, r2)
    assert "DUP_LOAN_PAYMENT_HISTORY" in c and "DUP_BANK_PAYMENT_HISTORY" in c


def test_single_fake_receipt_not_pooled(conn, tmp_path):
    rid, _, _ = _day(conn, tmp_path, 1, date(2026, 9, 28), 1000,
                     [("UPI/SBIN/626900001111/ABDUL RAHIM/UPI/x/a@ok", 0, 1000)],
                     [dict(vno=1, name="SUMA K", instr="4321", receipt=1000)])
    assert "TIMING_BOOK_RECEIPT" in codes(conn, rid)


def test_branch_without_totals_rejected(tmp_path):
    import openpyxl
    from app.parsers import parse_bank_book
    k = tmp_path / "k.xlsx"
    write_book(k, date(2026, 9, 28), 1000, [dict(vno=1, name="A", payment=100)])
    wb = openpyxl.load_workbook(k)
    wb.active.append([9, "Ghost Branch", None, date(2026, 9, 28), 777, "BP", "x", None,
                      "Direct Pay, A/c: PERS, A/c No: 1, Name: SURESH", "", 0, 40000, None, ""])
    wb.save(k)
    assert not parse_bank_book(k)["ok"]


def test_newest_first_statement_uses_dates(tmp_path):
    import openpyxl
    from app.parsers import parse_bank_statement
    p = tmp_path / "s.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["SlNo", "Transaction Date", "Value Date", "Particulars", None, "Cheque Number", "Withdrawals", None, "Deposits", "Balance Amount"])
    ws.append([2, "02-09-2026", "02-09-2026", "UPI/X/1/A/", None, None, None, None, "100.00", "1,100.00"])
    ws.append([1, "01-09-2026", "01-09-2026", "RAP/1/ 1 /B/", None, None, "100.00", None, None, "1,000.00"])
    wb.save(p)
    s = parse_bank_statement(p)
    assert s["ok"] and s["opening"] == 1100.0 and s["closing"] == 1100.0


def test_web_security(conn):
    from fastapi.testclient import TestClient
    from app.main import app, safe_back
    assert safe_back("//evil.com") == "/" and safe_back("/\\evil.com") == "/" and safe_back("/runs/1") == "/runs/1"
    c = TestClient(app, follow_redirects=False)
    r = c.get("/setup"); import re
    tok = re.search(r'name="csrf_token" value="([^"]+)"', r.text).group(1)
    r = c.post("/setup", data={"username": "kev", "password": "abcdefghijk1", "password2": "abcdefghijk1", "csrf_token": tok})
    assert c.get("/").status_code == 200
    stolen = dict(c.cookies)
    tok = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/").text).group(1)
    c.post("/logout", data={"csrf_token": tok})
    thief = TestClient(app, follow_redirects=False, cookies=stolen)
    assert thief.get("/").status_code == 303          # old cookie no longer works
    anon = TestClient(app, follow_redirects=False)
    assert anon.post("/reconcile/cancel", data={"token": "abc"}).status_code == 303
    assert anon.get("/runs/1").status_code == 303


def test_accounts_bulk_import(conn, tmp_path):
    import openpyxl
    from app.accounts_import import parse_accounts, apply_import
    p = tmp_path / "a.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append([None, "Aloor", "   SOUTH INDIAN BANK A/C NO 0058073000001317"])
    ws.append([None, "North Paravur", "SOUTH INDIAN BANK A/C NO 0058073000001113"])
    ws.append([None, "Head Office", "SOUTH INDIAN BANK A/C NO 0058073000001124"])
    ws.append([None, "Head Office", "SOUTH INDIAN BANK A/C NO 00580730001195"])
    wb.save(p)
    rows, problems = parse_accounts(str(p), main_last4={"1113"})
    assert len(rows) == 4 and {r["kind"] for r in rows if r["acc_no"].endswith("1113")} == {"main"}
    assert any("14 digits" in x for x in problems)
    r1 = apply_import(conn, rows, "t")
    assert r1["added"] == 4
    r2 = apply_import(conn, rows, "t")
    assert r2["added"] == 0 and r2["unchanged"] == 4


def test_new_narrations_and_sweeps(conn, tmp_path):
    """29-09-2026 SIB statement: sweeps, cash, RTGS, IMPS, IB all match a book booked by branch name."""
    src = SAMPLES / "SIB_1113_29-09-2026.xls"
    if not src.exists():
        pytest.skip("sample not present")
    from app.parsers import parse_bank_statement
    st = parse_bank_statement(src)
    branches = {"KALADY": "Kaladi", "AZHIKODE": "Azhicode", "KUTHIATHODE": "Kuthiyathode",
                "EAST KADUNGALLUR ALUVA": "East Kadungallur", "MUTTUCHIRA KADUTHURUTHY": "Muttuchira"}
    entries = []
    for n, r in enumerate(st["rows"], 1):
        if r["channel"] == "SWEEP":
            br = branches.get(r["branch"], r["branch"].title())
            entries.append(dict(vno=n, vtype="CT", name="", receipt=r["deposit"],
                                narration=f"Transfer from {br} collection account", ))
        elif r["deposit"]:
            entries.append(dict(vno=n, name=r["name"], receipt=r["deposit"]))
        else:
            entries.append(dict(vno=n, name=r["name"], payment=r["withdrawal"]))
    k = tmp_path / "k.xlsx"
    write_book(k, st["period_from"], st["opening"], entries)
    rid = run(conn, 1, src, k)
    b = brs(conn, rid)
    assert b["tallies"] and b["counts"] == {"a": 0, "b": 0, "c": 0, "d": 0}, b
    methods = [m["method"] for m in conn.execute("SELECT method FROM matches WHERE run_id=?", (rid,))]
    assert sum(m.startswith("Sweep from branch") for m in methods) == 33
    assert "RAP_RETURNED" in codes(conn, rid)


def _upi(ref, name, amt):
    return (f"UPI/SBIN/{ref}/{name}/UPI/x/p@ok", 0, amt)


def test_sunday_receipts_booked_on_monday(conn, tmp_path):
    sat, sun, mon = date(2026, 9, 26), date(2026, 9, 27), date(2026, 9, 28)
    r1, bc, kc = _day(conn, tmp_path, 1, sat, 100000,
                      [_upi("626900001234", "ANU", 500)],
                      [dict(vno=1, name="ANU", loan_ac="01", instr="1234", receipt=500)])
    # Sunday: three UPIs reach the bank; branches closed, ERP bank book has no entries
    r2, bc, kc2 = _day(conn, tmp_path, 2, sun, bc,
                       [_upi("626900001111", "ASHA MOL", 800), _upi("626900002222", "BINDU RAJ", 500),
                        _upi("626900003333", "CHINNU K", 700)], [], book_op=kc)
    assert kc2 == kc
    b = brs(conn, r2)
    assert b["tallies"] and b["c_holiday"] == 2000 and b["c_holiday_count"] == 3
    c = codes(conn, r2)
    assert "HOLIDAY_RECEIPTS_PENDING" in c and "BANK_CREDIT_NOT_IN_BOOK" not in c
    assert not [f for f in flags(conn, r2) if f["severity"] == "critical"]
    s = json.loads(conn.execute("SELECT summary_json FROM runs WHERE id=?", (r2,)).fetchone()[0])
    assert s["holidays"] == ["2026-09-27"]
    # Monday: ASHA booked with the Sunday instrument date, BINDU with Monday's date, CHINNU not booked
    r3, _, _ = _day(conn, tmp_path, 3, mon, bc, [_upi("626900004444", "DEVI S", 900)], book_op=kc, book_rows=[
        dict(vno=2, name="DEVI S", loan_ac="04", instr="4444", receipt=900),
        dict(vno=3, name="ASHA MOL", loan_ac="02", instr="1111", instr_date=sun, receipt=800),
        dict(vno=4, name="BINDU RAJ", loan_ac="03", instr="2222", receipt=500),
    ])
    b = brs(conn, r3)
    assert b["tallies"] and b["add_bank_credits_not_in_book"] == 700
    fl = flags(conn, r3)
    crit = [f for f in fl if f["code"] == "HOLIDAY_RECEIPT_NOT_BOOKED"]
    assert len(crit) == 1 and crit[0]["severity"] == "critical"
    assert "27-09-2026" in crit[0]["detail"] and "CHINNU" in crit[0]["detail"]
    booked = [f for f in fl if f["code"] == "HOLIDAY_RECEIPTS_BOOKED"]
    assert booked and booked[0]["severity"] == "medium" and "BINDU" in booked[0]["detail"]
    assert "HOLIDAY_RECEIPTS_PENDING" not in [f["code"] for f in fl]


def test_public_holiday_and_backdated_sunday_entry(conn, tmp_path):
    tue, wed, thu = date(2026, 9, 29), date(2026, 9, 30), date(2026, 10, 1)
    r1, bc, kc = _day(conn, tmp_path, 1, tue, 100000, [_upi("626900001234", "ANU", 500)],
                      [dict(vno=1, name="ANU", loan_ac="01", instr="1234", receipt=500)])
    # Wednesday: a bank holiday — nothing in the bank book, one UPI and one sweep-like debit arrive
    r2, bc, kc = _day(conn, tmp_path, 2, wed, bc,
                     [_upi("626900005555", "ELSY", 1200), ("TR FROM CHERAI BRANCH 58.73.1137", 300, 0)], [], book_op=kc)
    c = codes(conn, r2)
    assert "HOLIDAY_RECEIPTS_PENDING" in c and "HOLIDAY_DEBIT_PENDING" in c
    assert "BANK_DEBIT_NOT_IN_BOOK" not in c
    # Thursday: neither is booked -> both critical, citing the holiday date
    r3, _, _ = _day(conn, tmp_path, 3, thu, bc, [_upi("626900006666", "FEMY", 100)],
                    [dict(vno=2, name="FEMY", loan_ac="06", instr="6666", receipt=100)], book_op=kc)
    fl = {f["code"]: f for f in flags(conn, r3)}
    assert fl["HOLIDAY_RECEIPT_NOT_BOOKED"]["severity"] == "critical"
    assert "30-09-2026" in fl["HOLIDAY_RECEIPT_NOT_BOOKED"]["detail"]
    assert fl["HOLIDAY_DEBIT_NOT_BOOKED"]["severity"] == "critical"


def test_book_entry_dated_sunday_is_flagged(conn, tmp_path):
    sun = date(2026, 9, 27)
    rid, _, _ = _day(conn, tmp_path, 1, sun, 100000, [_upi("626900001111", "ASHA MOL", 800)],
                     [dict(vno=9, name="ASHA MOL", loan_ac="02", instr="1111", receipt=800)])
    f = [x for x in flags(conn, rid) if x["code"] == "BOOK_ENTRY_ON_SUNDAY"]
    assert f and f[0]["severity"] == "high" and "Vchr 9" in f[0]["detail"]


def test_holiday_run_pages_render(conn, tmp_path):
    import re
    from io import BytesIO
    import openpyxl
    from fastapi.testclient import TestClient
    from app.main import app
    sat, sun = date(2026, 9, 26), date(2026, 9, 27)
    _, bc, kc = _day(conn, tmp_path, 1, sat, 100000, [_upi("626900001234", "ANU", 500)],
                     [dict(vno=1, name="ANU", loan_ac="01", instr="1234", receipt=500)])
    r2, _, _ = _day(conn, tmp_path, 2, sun, bc, [_upi("626900001111", "ASHA MOL", 800)], [], book_op=kc)
    c = TestClient(app, follow_redirects=False)
    tok = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/setup").text).group(1)
    c.post("/setup", data={"username": "kev", "password": "abcdefghijk1", "password2": "abcdefghijk1", "csrf_token": tok})
    page = c.get(f"/runs/{r2}")
    assert page.status_code == 200 and "is a holiday" in page.text and "received on a holiday" in page.text
    x = c.get(f"/runs/{r2}/excel")
    assert x.status_code == 200
    wb = openpyxl.load_workbook(BytesIO(x.content))
    assert "holiday" in wb["BRS"]["A13"].value
    assert any("Holiday: received Sun 27-09-2026" == (r[10].value or "") for r in wb["Open Items"].iter_rows(min_row=4))


def test_listed_bank_holiday_backdated_entry(conn, tmp_path):
    tue, fri, sat = date(2026, 9, 29), date(2026, 10, 2), date(2026, 10, 3)
    conn.execute("INSERT INTO holidays(day,name,created_at) VALUES('2026-10-02','Gandhi Jayanthi','x')")
    _, bc, kc = _day(conn, tmp_path, 1, tue, 100000, [_upi("626900001234", "ANU", 500)],
                     [dict(vno=1, name="ANU", loan_ac="01", instr="1234", receipt=500)])
    # Gandhi Jayanthi: one UPI arrives; someone later posts a voucher dated that day for a different payer
    r2, bc, kc = _day(conn, tmp_path, 2, fri, bc, [_upi("626900002222", "BINDU RAJ", 600)],
                      [dict(vno=7, name="GEETHA", loan_ac="09", instr="9999", receipt=300)], book_op=kc,
                      acknowledged=True, ack_reason="Gap over 30-09 and 01-10, no transactions")
    fl = {f["code"]: f for f in flags(conn, r2)}
    assert fl["BOOK_ENTRY_ON_HOLIDAY"]["severity"] == "high" and "Gandhi Jayanthi" in fl["BOOK_ENTRY_ON_HOLIDAY"]["title"]
    assert "HOLIDAY_RECEIPTS_PENDING" in fl          # listed holiday even though the book has an entry
    r3, _, _ = _day(conn, tmp_path, 3, sat, bc, [_upi("626900003333", "CAROL", 100)],
                    [dict(vno=8, name="CAROL", loan_ac="10", instr="3333", receipt=100)], book_op=kc)
    fl = {f["code"]: f for f in flags(conn, r3)}
    assert fl["HOLIDAY_RECEIPT_NOT_BOOKED"]["severity"] == "critical"
    assert "Gandhi Jayanthi" in fl["HOLIDAY_RECEIPT_NOT_BOOKED"]["detail"]


def test_holidays_page(conn):
    import re
    from fastapi.testclient import TestClient
    from app.main import app, _parse_holiday_lines
    good, bad = _parse_holiday_lines("02-10-2026 Gandhi Jayanthi\n2026-12-25, Christmas\n2 Oct 2027 GJ\nnonsense line")
    assert [g[0].isoformat() for g in good] == ["2026-10-02", "2026-12-25", "2027-10-02"]
    assert good[0][1] == "Gandhi Jayanthi" and good[1][1] == "Christmas" and bad == ["nonsense line"]
    c = TestClient(app, follow_redirects=False)
    tok = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/setup").text).group(1)
    c.post("/setup", data={"username": "kev", "password": "abcdefghijk1", "password2": "abcdefghijk1", "csrf_token": tok})
    tok = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/holidays").text).group(1)
    assert c.post("/holidays", data={"lines": "25-12-2026 Christmas", "csrf_token": tok}).status_code == 303
    page = c.get("/holidays").text
    assert "<td>Christmas</td>" in page
    c.post("/holidays/delete", data={"day": "2026-12-25", "csrf_token": tok})
    assert "<td>Christmas</td>" not in c.get("/holidays").text
    assert TestClient(app, follow_redirects=False).get("/holidays").status_code == 303


def _client_setup():
    import re
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app, follow_redirects=False)
    tok = re.search(r'name="csrf_token" value="([^"]+)"', c.get("/setup").text).group(1)
    c.post("/setup", data={"username": "kev", "password": "abcdefghijk1", "password2": "abcdefghijk1", "csrf_token": tok})
    return c


def _tok(client, path="/"):
    import re
    r = client.get(path, follow_redirects=True)
    return re.search(r'name="csrf_token" value="([^"]+)"', r.text).group(1)


def _login(username, password):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app, follow_redirects=False)
    tok = _tok(c, "/login")
    r = c.post("/login", data={"username": username, "password": password, "csrf_token": tok})
    return c, r


def test_users_and_roles(conn):
    admin = _client_setup()
    t = _tok(admin, "/admin")
    for u, role in (("anjali", "operator"), ("vinod", "viewer"), ("meera", "manager")):
        r = admin.post("/admin/users", data={"username": u, "full_name": u.title(), "role": role,
                                             "password": "Temp-pass-123", "csrf_token": t})
        assert r.status_code == 303
    assert "Anjali" in admin.get("/admin").text
    # first sign-in forces a password change
    op, r = _login("anjali", "Temp-pass-123")
    assert r.headers["location"].endswith("/me")
    assert op.get("/reconcile").headers["location"].endswith("/me")
    t2 = _tok(op, "/me")
    op.post("/password", data={"current": "Temp-pass-123", "new": "anjali-own-pw1", "new2": "anjali-own-pw1", "csrf_token": t2})
    assert op.get("/reconcile").status_code == 200                       # operator may reconcile
    assert op.get("/admin").status_code == 303                           # ...but not the admin portal
    t2 = _tok(op, "/")
    op.post("/holidays", data={"lines": "25-12-2026 Christmas", "csrf_token": t2})
    assert conn.execute("SELECT COUNT(*) FROM holidays").fetchone()[0] == 1      # operators manage holidays
    op.post("/holidays/delete", data={"day": "2026-12-25", "csrf_token": t2})
    assert conn.execute("SELECT COUNT(*) FROM holidays").fetchone()[0] == 0
    assert op.get("/audit").status_code == 303
    # viewer: read-only
    vw, _ = _login("vinod", "Temp-pass-123")
    t3 = _tok(vw, "/me")
    vw.post("/password", data={"current": "Temp-pass-123", "new": "vinod-own-pw12", "new2": "vinod-own-pw12", "csrf_token": t3})
    assert vw.get("/").status_code == 200 and vw.get("/reconcile").status_code == 303
    assert "New reconciliation" not in vw.get("/").text
    # disabling a user signs them out at once
    uid = conn.execute("SELECT id FROM users WHERE username='anjali'").fetchone()[0]
    admin.post(f"/admin/users/{uid}", data={"full_name": "Anjali", "role": "operator", "active": "0", "csrf_token": t})
    assert op.get("/").headers["location"].endswith("/login")
    _, r = _login("anjali", "anjali-own-pw1")
    assert r.headers["location"].endswith("/login")
    # the only admin cannot demote or disable themselves
    kid = conn.execute("SELECT id FROM users WHERE username='kev'").fetchone()[0]
    admin.post(f"/admin/users/{kid}", data={"full_name": "Kev", "role": "viewer", "active": "0", "csrf_token": t})
    assert conn.execute("SELECT role, active FROM users WHERE id=?", (kid,)).fetchone()[:] == ("admin", 1)
    # password reset forces a new password and signs out
    vid = conn.execute("SELECT id FROM users WHERE username='vinod'").fetchone()[0]
    admin.post(f"/admin/users/{vid}/reset", data={"password": "Another-temp-1", "csrf_token": t})
    assert vw.get("/").status_code == 303
    assert conn.execute("SELECT must_change_pw FROM users WHERE id=?", (vid,)).fetchone()[0] == 1


def test_revoke_day_and_rerun(conn, tmp_path):
    d1, d2, d3 = date(2026, 9, 28), date(2026, 9, 29), date(2026, 9, 30)
    r1, bc1, kc1 = _day(conn, tmp_path, 1, d1, 100000, [_upi("626900001111", "ASHA", 100)],
                        [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100)])
    r2, bc2, kc2 = _day(conn, tmp_path, 2, d2, bc1, [_upi("626900002222", "BINU", 200)],
                        [dict(vno=2, name="BINU", loan_ac="2", instr="2222", receipt=200)])
    r3, _, _ = _day(conn, tmp_path, 3, d3, bc2, [_upi("626900003333", "CIJI", 300)],
                    [dict(vno=3, name="CIJI", loan_ac="3", instr="3333", receipt=300)])
    admin = _client_setup()
    assert "#" + str(r3) in admin.get(f"/admin/revoke?account=1&from={d2.isoformat()}").text
    t = _tok(admin, "/admin/revoke")
    r = admin.post("/admin/revoke", data={"account": "1", "from_date": d2.isoformat(), "reason": "short",
                                          "confirm": "REVOKE", "csrf_token": t})
    assert conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 3                 # reason too short
    admin.post("/admin/revoke", data={"account": "1", "from_date": d2.isoformat(),
                                      "reason": "Wrong bank book uploaded for 29-09", "confirm": "REVOKE", "csrf_token": t})
    assert [r[0] for r in conn.execute("SELECT id FROM runs ORDER BY id")] == [r1]
    arch = conn.execute("SELECT run_id, snapshot_json FROM revoked_runs ORDER BY id").fetchall()
    assert [a[0] for a in arch] == [r3, r2] and '"items"' in arch[0][1]
    assert "Wrong bank book" in admin.get("/admin/revoke").text
    # the same files can be reconciled again
    rr = run(conn, 1, tmp_path / "b2.xlsx", tmp_path / "k2.xlsx")
    assert brs(conn, rr)["tallies"]


def test_signoff_maker_checker(conn, tmp_path):
    rid, _, _ = _day(conn, tmp_path, 1, date(2026, 9, 28), 100000, [_upi("626900001111", "ASHA", 100)],
                     [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100)])
    admin = _client_setup()                      # "kev" — the run was made by "test", so kev may check it
    t = _tok(admin, f"/runs/{rid}")
    assert "Sign off as checked" in admin.get(f"/runs/{rid}").text
    admin.post(f"/runs/{rid}/signoff", data={"note": "Reviewed exceptions", "csrf_token": t})
    row = conn.execute("SELECT checked_by, checked_note FROM runs WHERE id=?", (rid,)).fetchone()
    assert row[0] == "kev" and row[1] == "Reviewed exceptions"
    conn.execute("UPDATE runs SET checked_by='', created_by='kev' WHERE id=?", (rid,))
    admin.post(f"/runs/{rid}/signoff", data={"csrf_token": t})
    assert conn.execute("SELECT checked_by FROM runs WHERE id=?", (rid,)).fetchone()[0] == ""   # can't check own run


def test_renewal_net_settlement(conn, tmp_path):
    d = date(2026, 9, 29)
    rid, _, _ = _day(conn, tmp_path, 1, d, 1_000_000, [
        ("RAP/626900000001/ 0000095001 /ANANDA PADMANABHAN/", 29500, 0),
        ("RAP/626900000002/ 0000095002 /MINI JOSE/", 10000, 0),
        ("UPI/SBIN/626900007777/MINI JOSE/UPI/x/p@ok", 0, 5000),
    ], [
        dict(vno=11, name="ANANDA PADMANABHAN", loan_ac="00110400700000000275", payment=180000),
        dict(vno=12, name="ANANDA PADMANABHAN", loan_ac="00110400700000000261", receipt=150500),
        # MINI's "closure" carries a real UPI ref that is on a bank credit: it is a collection, not a renewal
        dict(vno=13, name="MINI JOSE", loan_ac="0021", payment=15000),
        dict(vno=14, name="MINI JOSE", loan_ac="0022", instr="7777", receipt=5000),
    ])
    items = engine.run_items(conn, rid)
    m_of = {i["name"] + str(i["amount"]): i["match_id"] for i in items}
    mid = m_of["ANANDA PADMANABHAN29500.0"]
    assert mid and m_of["ANANDA PADMANABHAN180000.0"] == mid and m_of["ANANDA PADMANABHAN150500.0"] == mid
    row = conn.execute("SELECT method, confidence, note FROM matches WHERE id=?", (mid,)).fetchone()
    assert row[0] == "Renewal — net settled" and row[1] == "review" and "1,50,500.00" in row[2]
    f = [x for x in flags(conn, rid) if x["code"] == "RENEWAL_NETTED"]
    assert len(f) == 1 and f[0]["severity"] == "medium"
    assert m_of["MINI JOSE10000.0"] is None or m_of["MINI JOSE10000.0"] != m_of["MINI JOSE15000.0"]
    assert brs(conn, rid)["tallies"]


def test_short_branch_name_sweep(conn, tmp_path):
    d = date(2026, 9, 29)
    rid, _, _ = _day(conn, tmp_path, 1, d, 100000, [
        ("TR FROM KANICHUKULANGARA BRANCH 58.73.1137", 0, 45000),
    ], [dict(vno=1, name="", vtype="BR", receipt=45000, narration="Sweep received from Kanichu")])
    it = [i for i in engine.run_items(conn, rid) if i["side"] == "bank"][0]
    assert it["match_id"]


def test_exception_reviews_search_and_excel(conn, tmp_path):
    from io import BytesIO
    import openpyxl
    d1, d2 = date(2026, 9, 28), date(2026, 9, 29)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 100000, [_upi("626900001111", "ASHA", 100)],
                      [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100),
                       dict(vno=2, name="BINDU RAJAN", loan_ac="2", payment=30000)])
    admin = _client_setup()
    page = admin.get(f"/runs/{r1}").text
    assert 'id="flagsearch"' in page and "0 of" in page and "Unreviewed" in page
    fl = [f for f in flags(conn, r1) if f["code"] == "TIMING_BOOK_PAYMENT"][0]
    subject = engine.flag_subject(fl, r1)
    t = _tok(admin, f"/runs/{r1}")
    admin.post(f"/runs/{r1}/review", data={"subject": subject, "status": "followup", "remark": "", "csrf_token": t})
    assert conn.execute("SELECT COUNT(*) FROM reviews").fetchone()[0] == 0          # follow-up needs a remark
    admin.post(f"/runs/{r1}/review", data={"subject": subject, "status": "followup",
                                           "remark": "Branch to confirm RAP status", "csrf_token": t})
    page = admin.get(f"/runs/{r1}").text
    assert "Needs follow-up" in page and "Branch to confirm RAP status" in page
    # Excel: review on Exceptions sheet; level on Open Items
    wb = openpyxl.load_workbook(BytesIO(admin.get(f"/runs/{r1}/excel").content))
    ex = [[c.value for c in row] for row in wb["Exceptions"].iter_rows(min_row=2)]
    assert any(r[1] == "TIMING_BOOK_PAYMENT" and r[4] == "Needs follow-up" and r[6] == "kev" for r in ex)
    oi = wb["Open Items"]
    hdr = [c.value for c in oi[3]]
    assert hdr[11] == "Exception level"
    rows = [[c.value for c in row] for row in oi.iter_rows(min_row=4)]
    bindu = [r for r in rows if r[5] == "BINDU RAJAN"][0]
    assert bindu[11] == "MEDIUM" and "Needs follow-up" in bindu[13]
    # the review follows the item when it is carried to the next day
    r2, _, _ = _day(conn, tmp_path, 2, d2, bc, [_upi("626900002222", "CIJI", 200)],
                    [dict(vno=3, name="CIJI", loan_ac="3", instr="2222", receipt=200)], book_op=kc)
    assert "Branch to confirm RAP status" in admin.get(f"/runs/{r2}").text
    assert "BOOK_PAYMENT_NOT_IN_BANK" in codes(conn, r2)              # escalated, review kept


def test_recheck_applies_new_rules_to_old_runs(conn, tmp_path):
    d1, d2 = date(2026, 9, 28), date(2026, 9, 29)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 1_000_000, [
        ("RAP/626900000001/ 0000095001 /ANANDA PADMANABHAN/", 29500, 0),
        ("TR FROM KANICHUKULANGARA BRANCH 58.73.1137", 0, 45000),
        _upi("626900001111", "ASHA", 100),
    ], [
        dict(vno=11, name="ANANDA PADMANABHAN", loan_ac="0275", payment=180000),
        dict(vno=12, name="ANANDA PADMANABHAN", loan_ac="0261", receipt=150500),
        dict(vno=13, name="", vtype="BR", receipt=45000, narration="Sweep received from Kanichu"),
        dict(vno=14, name="ASHA", loan_ac="1", instr="1111", receipt=100),
    ])
    r2, _, _ = _day(conn, tmp_path, 2, d2, bc, [_upi("626900002222", "BINU", 200)],
                    [dict(vno=15, name="BINU", loan_ac="2", instr="2222", receipt=200)], book_op=kc)
    # pretend day 1 was reconciled before the new rules existed: drop those matches
    old = [r[0] for r in conn.execute("SELECT id FROM matches WHERE method LIKE 'Renewal%' OR method LIKE 'Sweep%'")]
    assert len(old) == 2
    for mid in old:
        conn.execute("UPDATE items SET match_id=NULL WHERE match_id=?", (mid,))
        conn.execute("DELETE FROM matches WHERE id=?", (mid,))
    engine.refresh_run(conn, r1); engine.refresh_run(conn, r2)
    assert "TIMING_BANK_DEBIT" in codes(conn, r1) and "BANK_DEBIT_NOT_IN_BOOK" in codes(conn, r2)
    manual_before = conn.execute("SELECT COUNT(*) FROM matches").fetchone()[0]

    out = engine.recheck_account(conn, 1, "test")
    assert out["new_matches"] == 2 and out["open_after"] == out["open_before"] - 5
    # recorded on day 1, so day 1's own exceptions are corrected too
    assert {r[0] for r in conn.execute("SELECT run_id FROM matches WHERE created_by='system (re-check)'")} == {r1}
    assert "TIMING_BANK_DEBIT" not in codes(conn, r1) and "RENEWAL_NETTED" in codes(conn, r1)
    assert "BANK_DEBIT_NOT_IN_BOOK" not in codes(conn, r2)
    assert conn.execute("SELECT COUNT(*) FROM matches").fetchone()[0] == manual_before + 2
    assert brs(conn, r1)["tallies"] and brs(conn, r2)["tallies"]

    # a match a person removes is never re-created automatically
    sweep = conn.execute("SELECT id FROM matches WHERE method LIKE 'Sweep%'").fetchone()[0]
    conn.execute("UPDATE matches SET run_id=? WHERE id=?", (r2, sweep))      # unmatch works on the latest run
    engine.unmatch(conn, sweep, "test", "wrong branch")
    assert engine.recheck_account(conn, 1, "test")["new_matches"] == 0


def test_mini_statement_grand_total_and_footer(tmp_path):
    import openpyxl
    from app.parsers import parse_bank_statement
    p = tmp_path / "mini.xlsx"
    d = date(2026, 10, 10)
    write_statement(p, d, 6000, [_upi("628338076055", "SAJNA K S", 980), _upi("628339674948", "ATHIRA SUNDARAN", 650),
                                 _upi("664912713456", "HARIKRISHNA A V", 850)])
    wb = openpyxl.load_workbook(p); ws = wb.active
    ws.delete_rows(ws.max_row)                                   # drop the "End of statement" line
    ws.append([None, "Grand Total:", None, None, None, None, "0", None, "2480", "8480.00 Cr"])
    ws.append([None, None, None, "This is a system-generated statement and does not require any signature"])
    ws.append(["Visit us at www.southindianbank.com. Customer Care 1800 425 1809"])
    ws.append(["This is a mini statement generated for viewing purpose only."])
    wb.save(p)
    out = parse_bank_statement(p)
    assert out["ok"], out["checks"]
    assert len(out["rows"]) == 3 and out["closing"] == 8480
    assert any(c["code"] == "BANK_GRAND_TOTAL" and c["level"] == "ok" for c in out["checks"])
    # a Grand Total that disagrees with the rows is an error
    ws = wb.active
    for row in ws.iter_rows():
        if row[1].value == "Grand Total:":
            row[8].value = "2400"
    wb.save(p)
    out = parse_bank_statement(p)
    assert not out["ok"] and any(c["code"] == "BANK_GRAND_TOTAL" for c in out["checks"])


def test_excel_bank_statement_has_customer_name(conn, tmp_path):
    from io import BytesIO
    import openpyxl
    from app import report
    rid, _, _ = _day(conn, tmp_path, 1, date(2026, 9, 29), 100000, [
        ("RAP/627311185948/ 0000095504 /VIJAYKRISHNA K R/", 37882, 0), ("DIVIN BIJU/", 1000, 0)],
        [dict(vno=1, name="VIJAYKRISHNA K R", loan_ac="1", payment=37882)])
    data, _ = report.build_report(conn, rid)
    ws = openpyxl.load_workbook(BytesIO(data))["Bank Statement"]
    assert [c.value for c in ws[1]][2:5] == ["Particulars", "Customer name", "Withdrawal"]
    names = [row[3].value for row in ws.iter_rows(min_row=2)]
    assert names[:2] == ["VIJAYKRISHNA K R", "DIVIN BIJU"]


def test_midnight_next_day_line_left_out_and_carried(conn, tmp_path):
    d1, d2 = date(2026, 10, 19), date(2026, 10, 20)
    mid = ("UPI/CNRB/629356446625/SARANYA S/UPI", 0, 200, d2)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 43530 - 300, [_upi("629259198205", "MANOHARAN K K", 300), mid],
                      [dict(vno=1, name="MANOHARAN K K", loan_ac="1", instr="8205", receipt=300)])
    run1 = conn.execute("SELECT * FROM runs WHERE id=?", (r1,)).fetchone()
    assert run1["bank_close"] == 43530 and brs(conn, r1)["tallies"]
    assert any(c["code"] == "BANK_NEXT_DAY" for c in json.loads(run1["checks_json"]))
    assert not [i for i in engine.run_items(conn, r1) if i["side"] == "bank" and i["date"] == d2]
    # next day's statement does NOT repeat the midnight line: it is brought in from yesterday's file
    r2, _, _ = _day(conn, tmp_path, 2, d2, 43730, [_upi("629300000001", "ANU", 100)],
                    [dict(vno=2, name="SARANYA S", loan_ac="2", instr="6625", receipt=200),
                     dict(vno=3, name="ANU", loan_ac="3", instr="0001", receipt=100)], book_op=kc)
    assert brs(conn, r2)["tallies"]
    assert not [i for i in engine.open_items_as_of(conn, 1, r2)]


def test_midnight_line_repeated_next_day(conn, tmp_path):
    d1, d2 = date(2026, 10, 19), date(2026, 10, 20)
    mid = ("UPI/CNRB/629356446625/SARANYA S/UPI", 0, 200, d2)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 1000, [_upi("629259198205", "MANO", 300), mid],
                      [dict(vno=1, name="MANO", loan_ac="1", instr="8205", receipt=300)])
    r2, _, _ = _day(conn, tmp_path, 2, d2, 1300, [mid[:3]],
                    [dict(vno=2, name="SARANYA S", loan_ac="2", instr="6625", receipt=200)], book_op=kc)
    assert brs(conn, r2)["tallies"]
    assert any(c["code"] == "BANK_NEXT_DAY_IN" for c in json.loads(conn.execute(
        "SELECT checks_json FROM runs WHERE id=?", (r2,)).fetchone()[0]))


def test_backdated_entries_found_from_fresh_book(conn, tmp_path):
    d1, d2 = date(2026, 10, 1), date(2026, 10, 2)
    # day 1: a big credit in the bank is not booked; a payment is booked but never debited
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 1_000_000, [("TR FROM CHERAI BRANCH 58.73.1137", 0, 1_292_382),
                                                          _upi("626900001111", "ASHA", 100)],
                      [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100),
                       dict(vno=2, name="WRONG PAYEE", loan_ac="9", payment=5000)])
    # later, in the ERP: the credit is booked dated day 1, and the wrong payment is deleted
    fresh = tmp_path / "fresh.xlsx"
    write_book(fresh, d1, 1_000_000, [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100),
                                      dict(vno=7, name="", vtype="BR", receipt=1_292_382,
                                           narration="Sweep received from Cherai")])
    new_open = 1_000_000 + 100 + 1_292_382
    b2, k2 = tmp_path / "b2.xlsx", tmp_path / "k2.xlsx"
    write_statement(b2, d2, bc, [_upi("626900002222", "BINU", 200)])
    write_book(k2, d2, new_open, [dict(vno=8, name="BINU", loan_ac="2", instr="2222", receipt=200)])
    with pytest.raises(engine.Blocked) as e:
        run(conn, 1, b2, k2)
    assert any(c["code"] == "BOOK_CONTINUITY" and c["level"] == "error" for c in e.value.checks)
    rid = engine.reconcile(conn, 1, str(b2), "b2", str(k2), "k2", "test", backdated_path=str(fresh),
                           backdated_name="fresh.xlsx", confirmed=True)
    b = brs(conn, rid)
    assert b["tallies"] and b["clean"]
    assert not engine.open_items_as_of(conn, 1, rid)          # sweep cleared by the late voucher; payment closed
    c = codes(conn, rid)
    assert "BOOK_BACKDATED_POSTED" in c and "BOOK_VOUCHER_REMOVED" in c
    checks = json.loads(conn.execute("SELECT checks_json FROM runs WHERE id=?", (rid,)).fetchone()[0])
    assert any(x["code"] == "CONFIRM_BACKDATED" and x["level"] == "ok" for x in checks)


def test_icici_csv_statement(tmp_path):
    from app.parsers import parse_bank_statement
    p = tmp_path / "icici.csv"
    p.write_text(
        'ICICI Bank account Statement from 19-10-2026 to 19-10-2026. \n\n\n'
        '"Account Number","Tran Date","Tran Particular","Tran Remarks","Inst Num","Orig Sol Id","Currency Code","Dr Tran Amt","Cr Tran Amt","Bal Amt","Deposit Branch"\n'
        '"--------------","---------","---------------","--------","-------","------------","-----------","-----------","-------","---------------"\n'
        '"Ac No:105305000993","19-OCT-2026","B/F","","","","INR","","","1000.00",""\n'
        '"Ac No:105305000993","19-Oct-2026","UPI/665852033974/UPI/XX03-1@okaxis/State Bank Of I /AXI85","665852033974//92","","2056","INR","0.00","300.00","1300.00","RPC KOTA"\n'
        '"Ac No:105305000993","19-Oct-2026","TPSL-19/10/2026-NACH5705-ICIC","","","0103","INR","0.00","200.00","1500.00","RPC MUMBAI"\n'
        '"Ac No:105305000993","19-Oct-2026","TO SIB 1113 SWEEP","","","0103","INR","500.00","0.00","1000.00","RPC MUMBAI"\n'
        '\n Closing Balance as on 20-10-2026 03:25:09 is INR.1000.00 includes Uncleared Funds of INR.0.00\n')
    out = parse_bank_statement(p)
    assert out["ok"], out["checks"]
    assert out["account_no"] == "105305000993" and out["opening"] == 1000 and out["closing"] == 1000
    r = out["rows"]
    assert (r[0]["channel"], r[0]["ref"]) == ("UPI", "665852033974") and r[1]["channel"] == "NACH"
    assert any(c["code"] == "BANK_BF" and c["level"] == "ok" for c in out["checks"])


def test_reconcile_form_requires_backdated_confirmation(conn, tmp_path):
    admin = _client_setup()
    t = _tok(admin, "/reconcile")
    b, k = tmp_path / "b.xlsx", tmp_path / "k.xlsx"
    write_statement(b, date(2026, 10, 1), 1000, [_upi("626900001111", "ASHA", 100)])
    write_book(k, date(2026, 10, 1), 1000, [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100)])
    files = {"bank_file": ("b.xlsx", b.read_bytes()), "book_file": ("k.xlsx", k.read_bytes())}
    r = admin.post("/reconcile", data={"account_id": "1", "csrf_token": t}, files=files)
    assert r.status_code == 303 and conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 0
    r = admin.post("/reconcile", data={"account_id": "1", "csrf_token": t, "confirm_backdated": "1"}, files=files)
    assert r.status_code == 200 and "confirm" in r.text.lower()      # first-run confirmation page


def test_bulk_upload_identifies_pairs_and_lists_missing(conn, tmp_path):
    import re
    conn.execute("UPDATE accounts SET acc_no='0058053000001113', branch='North Paravur' WHERE id=1")
    conn.execute("UPDATE accounts SET acc_no='0058073000001426', branch='Azhicode' WHERE id=2")
    for no in ("0058073000001124", "0058073000001177"):
        conn.execute("INSERT INTO accounts(bank,acc_no,nickname,kind,branch,created_at) VALUES('SIB',?,?,'collection','Head Office','x')",
                     (no, "HO " + no[-4:]))
    d = date(2026, 10, 1)
    f = {}
    def st(name, rows, op=1000):
        p = tmp_path / name; write_statement(p, d, op, rows); f[name] = p
    def bk(name, rows, branch, op=1000):
        p = tmp_path / name; write_book(p, d, op, rows, branch=branch); f[name] = p
    st("AccountStatement_XXXXXXXXXXXX1113.xlsx", [_upi("626900001111", "ASHA", 100)])
    bk("BankBook (1).xlsx", [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100)], "North Paravur")
    st("AccountStatement_XXXXXXXXXXXX1426.xlsx", [_upi("626900002222", "BINU", 200)])
    bk("BankBook (2).xlsx", [dict(vno=2, name="BINU", loan_ac="2", instr="2222", receipt=200)], "Azhicode")
    st("AccountStatement_XXXXXXXXXXXX1124.xlsx", [_upi("626900003333", "CIJI", 300)])
    bk("BankBook (3).xlsx", [dict(vno=3, name="CIJI", loan_ac="3", instr="3333", receipt=300)], "Head Office")
    (tmp_path / "notes.csv").write_text("hello,world\n")
    f["notes.csv"] = tmp_path / "notes.csv"
    admin = _client_setup()
    t = _tok(admin, "/reconcile")
    files = [("files", (n, p.read_bytes())) for n, p in f.items()]
    r = admin.post("/reconcile/batch", data={"csrf_token": t, "confirm_backdated": "1"}, files=files)
    bid = r.headers["location"].rsplit("/", 1)[-1]
    page = admin.get(f"/reconcile/batch/{bid}").text
    # first runs of new accounts need a confirmation; Head Office book is ambiguous (two accounts)
    assert page.count("Needs confirmation") == 2
    assert "Branch Head Office has" in page and "Choose which one" in page
    assert "notes.csv" in page and "be used" in page
    assert "HO 1124" in page and "Bank book" in page          # 1124 statement waits for its book
    assert "HO 1177" in page                                    # nothing uploaded for 1177
    # assign the Head Office book to 1124 -> paired and run
    from app.main import _batch_load
    b = _batch_load(bid)
    ho = [x for x in b["files"] if x["name"] == "BankBook (3).xlsx"][0]
    admin.post(f"/reconcile/batch/{bid}/assign", data={"csrf_token": t, "fid": ho["fid"], "account_id": "3"})
    b = _batch_load(bid)
    assert sum(1 for v in b["results"].values() if v["status"] == "confirm") == 3
    # confirm one first run through the usual page -> back to the bulk page, marked reconciled
    tok = next(v["token"] for v in b["results"].values() if v["status"] == "confirm")
    blocked = admin.get(f"/reconcile/resume/{tok}")
    assert blocked.status_code == 200 and "Please confirm before continuing" in blocked.text
    r = admin.post("/reconcile/ack", data={"csrf_token": t, "token": tok, "reason": "Confirmed account for first run"})
    assert r.headers["location"].endswith(f"/reconcile/batch/{bid}")
    b = _batch_load(bid)
    assert sum(1 for v in b["results"].values() if v["status"] == "done") == 1
    assert conn.execute("SELECT COUNT(*) FROM runs").fetchone()[0] == 1


def test_two_bank_books_nach_settlement_and_txn_id(conn, tmp_path):
    """ICICI: UPI book + NACH book read as one ledger; NACH credit = NACH book; wrong last-4 flagged."""
    d = date(2026, 10, 19)
    b, k1, k2 = tmp_path / "b.xlsx", tmp_path / "upi.xlsx", tmp_path / "nach.xlsx"
    write_statement(b, d, 1000, [_upi("626900001111", "ASHA", 300), _upi("626900002222", "BINU", 250),
                                 ("TPSL-19/10/2026-NACH5705-ICIC", 0, 700)])
    write_book(k1, d, 1500, [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=300),
                             dict(vno=2, name="BINU", loan_ac="2", instr="9999", receipt=250)], branch="Head Office")
    write_book(k2, d, -500, [dict(vno=3, name="CHITRA", loan_ac="3", instr="79820261017004825", receipt=400),
                             dict(vno=4, name="DEEPA", loan_ac="4", instr="80620261017004825", receipt=300)],
               branch="Head Office")
    rid = engine.reconcile(conn, 1, str(b), "b.xlsx", [str(k1), str(k2)], "upi.xlsx + nach.xlsx", "t",
                           acknowledged=True, ack_reason="first run for this account", confirmed=True)
    r = conn.execute("SELECT * FROM runs WHERE id=?", (rid,)).fetchone()
    assert r["book_open"] == 1000 and r["book_file"] == "upi.xlsx + nach.xlsx"
    assert json.loads(r["brs_json"])["clean"]
    methods = [m[0] for m in conn.execute("SELECT method FROM matches WHERE run_id=?", (rid,))]
    assert any(m.startswith("NACH settlement") for m in methods)
    assert conn.execute("SELECT COUNT(*) FROM items WHERE run_id=? AND match_id IS NULL", (rid,)).fetchone()[0] == 0
    codes = [f["code"] for f in json.loads(r["flags_json"])]
    assert codes.count("TXN_ID_MISMATCH") == 1 and "TXN_ID_CHECK" in codes
    items = [i for i in engine.run_items(conn, rid) if i["side"] == "book"]
    assert {i["book_name"] for i in items} == {"upi.xlsx", "nach.xlsx"}


def test_bulk_pairs_two_books_of_one_account(conn, tmp_path):
    from app import batch
    conn.execute("UPDATE accounts SET acc_no='105305000993', branch='Head Office', bank='ICICI' WHERE id=1")
    conn.execute("UPDATE accounts SET branch='Head Office' WHERE id=2")
    d = date(2026, 10, 19)
    b, k1, k2 = tmp_path / "b_105305000993.xlsx", tmp_path / "upi.xlsx", tmp_path / "nach.xlsx"
    write_statement(b, d, 1000, [_upi("626900001111", "ASHA", 300), ("TPSL-19/10/2026-NACH5705-ICIC", 0, 700)])
    write_book(k1, d, 1500, [dict(vno=1, name="ASHA", loan_ac="1", instr="626900001111", receipt=300)],
               branch="Head Office")
    write_book(k2, d, -500, [dict(vno=3, name="CHITRA", loan_ac="3", instr="79820261017004825", receipt=700)],
               branch="Head Office")
    accounts = [dict(a) for a in conn.execute("SELECT * FROM accounts")]
    files = []
    for i, p in enumerate((b, k1, k2)):
        info = batch.classify(str(p), p.name)
        aid, how = batch.identify(conn, info, accounts)
        files.append(dict(info, fid=f"f{i}", name=p.name, path=str(p), account_id=aid, how=how,
                          date_from=info["date_from"].isoformat() if info["date_from"] else None))
    assert files[0]["account_id"] == 1 and files[1]["account_id"] is None and files[2]["account_id"] is None
    batch.resolve_by_content(files)
    assert files[1]["account_id"] == 1 and "transaction ID" in files[1]["how"]
    assert files[2]["account_id"] == 1 and "NACH" in files[2]["how"]
    pl = batch.plan(files)
    assert len(pl["pairs"]) == 1 and pl["pairs"][0]["books"] == ["f1", "f2"] and not pl["lonely"]


def test_reconcile_form_accepts_two_bank_books(conn, tmp_path):
    admin = _client_setup()
    t = _tok(admin, "/reconcile")
    d = date(2026, 10, 19)
    b, k1, k2 = tmp_path / "b.xlsx", tmp_path / "upi.xlsx", tmp_path / "nach.xlsx"
    write_statement(b, d, 1000, [_upi("626900001111", "ASHA", 300), ("TPSL-19/10/2026-NACH5705-ICIC", 0, 700)])
    write_book(k1, d, 1500, [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=300)], branch="Head Office")
    write_book(k2, d, -500, [dict(vno=3, name="CHITRA", loan_ac="3", instr="79820261017004825", receipt=700)],
               branch="Head Office")
    files = [("bank_file", ("b.xlsx", b.read_bytes())), ("book_file", ("upi.xlsx", k1.read_bytes())),
             ("book_file", ("nach.xlsx", k2.read_bytes()))]
    r = admin.post("/reconcile", data={"account_id": "1", "csrf_token": t, "confirm_backdated": "1"}, files=files)
    assert r.status_code == 200 and "upi.xlsx + nach.xlsx" in r.text
    assert "2 bank books read together" in r.text


def test_large_own_transfer_and_name_in_narration_pair(conn, tmp_path):
    """29-09-2026: ₹8,10,30,000 from the company's own account (book JV 'AMOUNT TRANSFERED TO SIB 1113') stayed open."""
    d = date(2026, 9, 29)
    rid, _, _ = _day(conn, tmp_path, 1, d, 1000000,
                     [("FREEMEN INVESTMENTS PRIVATE LIMITED", 0, 81030000), _upi("626900001111", "ASHA", 740),
                      ("RTGS TO:SIBLH26272448778:K O VARGHESE", 70000000, 0)],
                     [dict(vno=1, vtype="JV", name="", receipt=81030000,
                           narration="AMOUNT TRANSFERED TO SIB 1113 81030000/-, A/c: IBV"),
                      dict(vno=2, name="BEENA", loan_ac="2", instr="9999", receipt=800),
                      dict(vno=3, name="", payment=70000000, narration="PART OF THE LOAN AMOUNT PAID TO VARGHESE K O")])
    got = {m[0]: m[1] for m in conn.execute("SELECT method, confidence FROM matches WHERE run_id=?", (rid,))}
    assert got.get("Own-account transfer (unique on the day)") == "strong"
    assert got.get("Amount + name in narration (unique on the day)") == "strong"


def test_entries_booked_late_with_earlier_date_adjust_without_manager(conn, tmp_path):
    """Bank lines of day 1 booked in the ERP on day 2 but dated day 1: the day-2 book opening moves by
    exactly those lines. The operator's run goes through (no acknowledgement) and the lines are closed."""
    d1, d2 = date(2026, 9, 28), date(2026, 9, 29)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 100000,
                      [("CHRG/SMS CHARGES Q3", 59, 0), ("NACH DR/ACH/LIC PREMIUM", 2000, 0),
                       _upi("626900001111", "ASHA", 300)],
                      [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=300)])
    assert "TIMING_BANK_DEBIT" in codes(conn, r1)
    r2, _, _ = _day(conn, tmp_path, 2, d2, bc, [_upi("626900002222", "BINU", 200)],
                    [dict(vno=3, name="BINU", loan_ac="2", instr="2222", receipt=200)],
                    book_op=round(kc - 2059, 2))
    assert brs(conn, r2)["clean"]
    assert "BOOK_LATE_POSTED" in codes(conn, r2) and "BANK_DEBIT_NOT_IN_BOOK" not in codes(conn, r2)
    m = conn.execute("SELECT method, run_id FROM matches WHERE method LIKE 'Booked late%'").fetchone()
    assert m["run_id"] == r2


def test_inbox_alerts_managers_not_operators(conn, tmp_path):
    admin = _client_setup()
    t = _tok(admin, "/admin")
    admin.post("/admin/users", data={"username": "anjali", "full_name": "Anjali", "role": "operator",
                                     "password": "Temp-pass-123", "csrf_token": t})
    op, _ = _login("anjali", "Temp-pass-123")
    t2 = _tok(op, "/me")
    op.post("/password", data={"current": "Temp-pass-123", "new": "anjali-own-pw1", "new2": "anjali-own-pw1",
                               "csrf_token": t2})
    d1, d2 = date(2026, 9, 28), date(2026, 9, 29)
    r1, bc, kc = _day(conn, tmp_path, 1, d1, 100000, [("CHRG/SMS CHARGES Q3", 59, 0), _upi("626900001111", "ASHA", 300)],
                      [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=300)])
    r2, _, _ = _day(conn, tmp_path, 2, d2, bc, [_upi("626900002222", "BINU", 200)],
                    [dict(vno=3, name="BINU", loan_ac="2", instr="2222", receipt=200)], book_op=round(kc - 59, 2))
    titles = [r[0] for r in conn.execute("SELECT title FROM alerts WHERE title NOT LIKE '%confirmed by%' ORDER BY id")]
    assert len(titles) == 2 and "timing difference" in titles[0] and "booked late" in titles[1]
    assert op.get("/inbox").status_code == 303 and 'class="dot"' not in op.get("/").text   # operators: no inbox
    page = admin.get("/inbox").text
    assert "booked late" in page and "Mark all read" in page and 'class="dot">3<' in admin.get("/").text
    admin.post("/inbox/read", data={"alert_id": "all", "csrf_token": _tok(admin, "/inbox")})
    assert 'class="dot"' not in admin.get("/").text
    engine.undo_run(conn, r2, "kev")
    assert conn.execute("SELECT COUNT(*) FROM alerts").fetchone()[0] == 2           # day 1 alerts stay


def test_operator_confirms_first_day_and_managers_are_alerted(conn, tmp_path):
    admin = _client_setup()
    t = _tok(admin, "/admin")
    admin.post("/admin/users", data={"username": "neethu", "full_name": "Neethu", "role": "operator",
                                     "password": "Temp-pass-123", "csrf_token": t})
    op, _ = _login("neethu", "Temp-pass-123")
    op.post("/password", data={"current": "Temp-pass-123", "new": "neethu-own-pw1", "new2": "neethu-own-pw1",
                               "csrf_token": _tok(op, "/me")})
    d = date(2026, 10, 19)
    b, k = tmp_path / "b.xlsx", tmp_path / "k.xlsx"
    write_statement(b, d, 1000, [_upi("626900001111", "ASHA", 100)])
    write_book(k, d, 1000, [dict(vno=1, name="ASHA", loan_ac="1", instr="1111", receipt=100)])
    t = _tok(op, "/reconcile")
    r = op.post("/reconcile", data={"account_id": "1", "csrf_token": t, "confirm_backdated": "1"},
                files={"bank_file": ("b.xlsx", b.read_bytes()), "book_file": ("k.xlsx", k.read_bytes())})
    assert "Proceed anyway?" in r.text and "A Manager or Admin must confirm" not in r.text
    import re as _re
    token = _re.search(r'name="token" value="([^"]+)"', r.text).group(1)
    r = op.post("/reconcile/ack", data={"token": token, "reason": "Aloor mini statement, starting point agreed",
                                        "csrf_token": _tok(op, "/reconcile")})
    assert r.status_code == 303 and "/runs/" in r.headers["location"]
    al = conn.execute("SELECT title, body FROM alerts").fetchall()
    assert any("confirmed by neethu" in a[0] and "starting point agreed" in a[1] for a in al)


def test_first_day_with_matching_account_number_needs_no_confirmation(conn, tmp_path):
    from app.parsers import parse_bank_statement
    S = "/tmp/claude-0/-home-claude/0d0f0921-aaea-52b2-8732-a729caf91e89/scratchpad/in29/"
    import os
    if not os.path.exists(S + "stmt.xls"):
        pytest.skip("real 29-09 statement not available")
    conn.execute("UPDATE accounts SET acc_no='0058073000001113' WHERE id=1")
    try:
        engine.reconcile(conn, 1, S + "stmt.xls", "stmt.xls", S + "book.xlsx", "book.xlsx", "op", confirmed=True)
    except engine.Blocked as e:
        assert "FIRST_RUN" not in {c["code"] for c in e.ackable}       # only the opening difference remains
